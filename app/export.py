"""TXT, SRT, DOCX ve PDF dışa aktarma. PDF'te yazı tipi gömülür; Arapça, Çince ve Japonca da desteklenir."""

import io
import re

from . import fonts

PT_TO_MM = 0.3528


def clock(sec: float, srt: bool = False) -> str:
    ms = int(round(sec * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    if srt:
        return f"{h:02}:{m:02}:{s:02},{ms:03}"
    return f"{h}:{m:02}:{s:02}" if h else f"{m:02}:{s:02}"


def raw_blocks(segments: list[dict], timestamps: bool) -> list[tuple[str, str]]:
    """Ham transkripti (tür, metin) bloklarına çevirir; uzun duraklarda paragraf açar."""
    if timestamps:
        return [("p", f"[{clock(s['start'])}] {s['text']}") for s in segments]
    blocks, cur, last_end = [], [], None
    for s in segments:
        if cur and (s["start"] - last_end > 1.5 or len(cur) >= 6):
            blocks.append(("p", " ".join(cur)))
            cur = []
        cur.append(s["text"])
        last_end = s["end"]
    if cur:
        blocks.append(("p", " ".join(cur)))
    return blocks


def markdown_blocks(text: str) -> list[tuple[str, str]]:
    """Basit Markdown: #/##/### başlık, '- ' madde, boş satırla ayrılan paragraflar."""
    blocks, para = [], []

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for line in text.splitlines():
        line = line.strip()
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if not line:
            flush()
        elif m:
            flush()
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip("* ")))
        elif re.match(r"^([-*•]|\d+[.)])\s+", line):
            flush()
            blocks.append(("li", re.sub(r"^([-*•])\s+", "", line)))
        else:
            para.append(line)
    flush()
    return blocks


def to_txt(blocks, title: str | None, date_text: str) -> bytes:
    out = []
    if title:
        out += [title, date_text, ""]
    for kind, text in blocks:
        plain = text.replace("**", "")
        out.append(f"• {plain}" if kind == "li" else plain)
        out.append("")
    return "\n".join(out).strip().encode("utf-8") + b"\n"


def to_srt(segments: list[dict]) -> bytes:
    lines = []
    for i, s in enumerate(segments, 1):
        lines += [str(i), f"{clock(s['start'], True)} --> {clock(s['end'], True)}", s["text"], ""]
    return "\n".join(lines).encode("utf-8")


def to_docx(blocks, title: str | None, date_text: str, font_id: str, size: int, spacing: float) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    family = fonts.FONTS[font_id]
    doc = Document()

    def style_font(style, pt, bold=None):
        style.font.name = family
        style.font.size = Pt(pt)
        if bold is not None:
            style.font.bold = bold
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        # Word, ana yazı tipinde olmayan alfabeleri (Arapça, Çince…) kendi yedekleriyle gösterir
        for attr in ("w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), family)

    normal = doc.styles["Normal"]
    style_font(normal, size)
    normal.paragraph_format.line_spacing = spacing
    normal.paragraph_format.space_after = Pt(size * 0.6)
    for level, scale in ((1, 1.5), (2, 1.25), (3, 1.1)):
        hs = doc.styles[f"Heading {level}"]
        style_font(hs, round(size * scale), bold=True)
        hs.font.color.rgb = None

    def finish(par, text):
        if fonts.is_rtl(text):
            par.paragraph_format.element.get_or_add_pPr().append(OxmlElement("w:bidi"))
        return par

    def add_runs(par, text):
        for i, piece in enumerate(text.split("**")):
            if piece:
                par.add_run(piece).bold = i % 2 == 1
        finish(par, text)

    if title:
        finish(doc.add_heading(title, level=1), title)
        if date_text:
            doc.add_paragraph(date_text)
    for kind, text in blocks:
        if kind.startswith("h"):
            finish(doc.add_heading(text, level=min(int(kind[1]) + (1 if title else 0), 3)), text)
        elif kind == "li":
            add_runs(doc.add_paragraph(style="List Bullet"), text)
        else:
            add_runs(doc.add_paragraph(), text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(blocks, title: str | None, date_text: str, font_id: str, size: int, spacing: float) -> bytes:
    from fpdf import FPDF

    fonts.ensure(font_id)
    pdf = FPDF(format="A4")
    pdf.set_margins(22, 22, 22)
    pdf.set_auto_page_break(True, margin=22)
    pdf.add_font("Govde", "", str(fonts.path(font_id, False)))
    pdf.add_font("Govde", "B", str(fonts.path(font_id, True)))

    all_text = " ".join([title or ""] + [t for _, t in blocks])
    extra = fonts.fallbacks_for(all_text)
    for fid, fpath in extra:
        pdf.add_font(fid, "", fpath)
    if extra:
        pdf.set_fallback_fonts([fid for fid, _ in extra], exact_match=False)
        # Arapça harf birleştirme ve sağdan sola sıralama için metin şekillendirme
        pdf.set_text_shaping(True)
    pdf.add_page()

    def write(text, pt, bold=False, after=0.6, indent=0):
        pdf.set_font("Govde", "B" if bold else "", pt)
        pdf.set_x(pdf.l_margin + indent)
        # fpdf2 markdown: sadece **kalın** kullanılsın diye diğer işaretleri etkisizleştir
        safe = text.replace("--", "–").replace("__", "_")
        pdf.multi_cell(0, pt * spacing * PT_TO_MM, safe, align="R" if fonts.is_rtl(text) else "L",
                       markdown="**" in safe and not bold, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(pt * after * PT_TO_MM)

    if title:
        write(title, size * 1.7, bold=True, after=0.2)
        if date_text:
            pdf.set_text_color(92, 91, 86)
            write(date_text, size * 0.85, after=1.2)
            pdf.set_text_color(26, 26, 24)
    for kind, text in blocks:
        if kind.startswith("h"):
            scale = {"h1": 1.45, "h2": 1.25, "h3": 1.1}[kind]
            pdf.ln(size * 0.4 * PT_TO_MM)
            write(text, size * scale, bold=True, after=0.4)
        elif kind == "li":
            write(f"•  {text}", size, after=0.3, indent=4)
        else:
            write(text, size)
    return bytes(pdf.output())


def build(fmt: str, *, text: str | None, segments: list[dict] | None, use_raw: bool, title: str | None,
          date_text: str, timestamps: bool, font_id: str, size: int, spacing: float) -> bytes:
    if fmt == "srt":
        return to_srt(segments or [])
    blocks = raw_blocks(segments or [], timestamps) if use_raw else markdown_blocks(text or "")
    if fmt == "txt":
        return to_txt(blocks, title, date_text)
    if fmt == "docx":
        return to_docx(blocks, title, date_text, font_id, size, spacing)
    if fmt == "pdf":
        return to_pdf(blocks, title, date_text, font_id, size, spacing)
    raise ValueError(fmt)
